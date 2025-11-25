#!/usr/bin/env python3
"""
Extract content from a .docx file and convert to markdown.
"""

import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Error: python-docx is not installed. Install it with: pip install python-docx")
    sys.exit(1)

def docx_to_markdown(docx_path: str) -> str:
    """Convert a .docx file to markdown format."""
    doc = Document(docx_path)
    markdown_lines = []
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            markdown_lines.append("")
            continue
        
        # Check paragraph style for headings
        style_name = paragraph.style.name.lower()
        
        if 'heading 1' in style_name or 'title' in style_name:
            markdown_lines.append(f"# {text}")
        elif 'heading 2' in style_name:
            markdown_lines.append(f"## {text}")
        elif 'heading 3' in style_name:
            markdown_lines.append(f"### {text}")
        elif 'heading 4' in style_name:
            markdown_lines.append(f"#### {text}")
        elif 'heading 5' in style_name:
            markdown_lines.append(f"##### {text}")
        elif 'heading 6' in style_name:
            markdown_lines.append(f"###### {text}")
        else:
            # Regular paragraph - preserve formatting
            formatted_text = ""
            for run in paragraph.runs:
                run_text = run.text
                if run.bold:
                    run_text = f"**{run_text}**"
                if run.italic:
                    run_text = f"*{run_text}*"
                formatted_text += run_text
            
            if formatted_text:
                markdown_lines.append(formatted_text)
            else:
                markdown_lines.append(text)
    
    # Process tables
    for table in doc.tables:
        markdown_lines.append("")  # Add spacing before table
        
        # Process header row
        if table.rows:
            header_row = table.rows[0]
            headers = [cell.text.strip().replace('\n', ' ') for cell in header_row.cells]
            markdown_lines.append("| " + " | ".join(headers) + " |")
            markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            
            # Process data rows
            for row in table.rows[1:]:
                cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                markdown_lines.append("| " + " | ".join(cells) + " |")
        
        markdown_lines.append("")  # Add spacing after table
    
    return "\n".join(markdown_lines)

if __name__ == "__main__":
    docx_file = "islr-app/docs/infrastructure/Open Health - iSLR app production handover notes for Triradial.docx"
    output_file = "islr-app/docs/infrastructure/Open Health - iSLR app production handover notes for Triradial.md"
    
    try:
        markdown_content = docx_to_markdown(docx_file)
        
        # Write to markdown file
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(markdown_content)
        
        print(f"Successfully converted {docx_file}")
        print(f"Output saved to: {output_file}")
        
        # Also print to stdout so you can see it
        print("\n" + "="*80)
        print("MARKDOWN CONTENT:")
        print("="*80)
        print(markdown_content)
        
    except Exception as e:
        print(f"Error processing file: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)